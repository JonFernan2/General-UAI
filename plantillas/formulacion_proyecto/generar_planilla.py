"""Genera la planilla Excel para la Entrega N°1 de Formulación de Proyecto de Título
(Modalidad Licitación), según lo exigido en la "Guía de Desarrollo Asignatura Proyecto
de Título" (UVM): Listado de Actividades, Cubicaciones, Listado de Materiales y
Cotizaciones (materiales/herramientas/máquinas/equipos).

Uso:
    python3 generar_planilla.py [ruta_salida.xlsx]
"""
from __future__ import annotations

import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AZUL = "1F4E78"
AZUL_CLARO = "D9E2F3"
AMARILLO = "FFF2CC"
GRIS_EJEMPLO = "808080"
BLANCO = "FFFFFF"

FUENTE_TITULO = Font(bold=True, size=14, color=AZUL)
FUENTE_SUBTITULO = Font(bold=True, size=11, color=AZUL)
FUENTE_ENCABEZADO = Font(bold=True, color=BLANCO)
FUENTE_EJEMPLO = Font(italic=True, color=GRIS_EJEMPLO)
RELLENO_ENCABEZADO = PatternFill("solid", fgColor=AZUL)
RELLENO_TITULO = PatternFill("solid", fgColor=AZUL_CLARO)
BORDE_FINO = Border(*(Side(style="thin", color="B7B7B7") for _ in range(4)))

UNIDADES = "ml,m2,m3,kg,un,pm,gl"
TIPOS_MATERIAL = "Material,Herramienta,Máquina,Equipo"

FILAS_TRABAJO = 40  # filas en blanco listas para usar en cada listado
FILAS_COTIZACION = 30


def _encabezado(ws, fila: int, columnas: list[str], anchos: list[int]) -> None:
    for idx, (texto, ancho) in enumerate(zip(columnas, anchos), start=1):
        celda = ws.cell(row=fila, column=idx, value=texto)
        celda.font = FUENTE_ENCABEZADO
        celda.fill = RELLENO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = BORDE_FINO
        ws.column_dimensions[get_column_letter(idx)].width = ancho
    ws.row_dimensions[fila].height = 30
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def _titulo(ws, texto: str, ncols: int, fila: int = 1) -> None:
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ncols)
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.font = FUENTE_TITULO
    celda.fill = RELLENO_TITULO
    celda.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[fila].height = 26


def _bordear_area(ws, fila_inicio: int, fila_fin: int, col_inicio: int, col_fin: int) -> None:
    for fila in range(fila_inicio, fila_fin + 1):
        for col in range(col_inicio, col_fin + 1):
            ws.cell(row=fila, column=col).border = BORDE_FINO


def _fila_ejemplo(ws, fila: int, valores: list) -> None:
    for idx, valor in enumerate(valores, start=1):
        celda = ws.cell(row=fila, column=idx, value=valor)
        celda.font = FUENTE_EJEMPLO


def hoja_instrucciones(wb) -> None:
    ws = wb.active
    ws.title = "Instrucciones"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110

    _titulo(ws, "PLANILLA FORMULACIÓN DE PROYECTO DE TÍTULO — ENTREGA N°1 (Modalidad Licitación)", 2)
    ws["A2"] = "Proyecto:"
    ws["A2"].font = FUENTE_SUBTITULO
    ws["B2"] = ""  # celda para que el estudiante escriba el nombre del proyecto
    ws["B2"].font = Font(bold=True, size=12)
    ws["B2"].fill = PatternFill("solid", fgColor="FFFFFF")
    ws["B2"].border = Border(bottom=Side(style="thin", color=AZUL))
    ws["A3"].value = "(nombre del proyecto ►)"
    ws["A3"].font = Font(italic=True, color=GRIS_EJEMPLO, size=9)

    filas = [
        "",
        "Esta planilla cubre los entregables en formato Excel de la Entrega N°1 de "
        "'Formulación de Proyecto de Título' (Modalidad Licitación): Listado de Actividades, "
        "Cubicaciones, Listado de Materiales y Cotizaciones. Las Especificaciones Técnicas "
        "modificadas se entregan aparte, en formato Word.",
        "",
        "REGLAS CLAVE A RESPETAR (extraídas de la pauta de la asignatura):",
        "",
        "1) Especificaciones Técnicas (entrega en Word, no en esta planilla)",
        "   • Modificar/complementar las EETT de Arquitectura y Sanitarias (edificación).",
        "   • Resaltar en AMARILLO toda modificación agregada y TACHAR toda eliminación.",
        "   • Detallar calidad, materialidad, tipo, formato, color, modelo y marca/proveedor "
        "de cada material, sin dejar nada a interpretación.",
        "",
        "2) Listado de Actividades (hoja 'Listado de Actividades')",
        "   • El ITEM debe coincidir exactamente con la numeración usada en las EETT.",
        "   • La descripción debe incluir características relevantes (medidas, espesores, dimensiones).",
        "",
        "3) Cubicaciones (hoja 'Cubicaciones')",
        "   • Cantidades EXACTAS, con al menos 2 decimales. NUNCA redondear al entero superior.",
        "   • Partidas en unidad 'un' deben cubicarse en números enteros (no se fraccionan).",
        "   • Cada cubicación debe incluir la memoria de cálculo (fórmula/desarrollo) que respalda el número.",
        "   • En Edificación se excluyen de evaluación: Instalaciones Eléctricas/CCDD/CCTV, Clima, "
        "Ascensores y Redes de Gas (igual se recomienda cubicarlas). Sanitarias, Aguas Lluvias, "
        "Alcantarillado, Pavimentación y Urbanización SÍ deben analizarse completas.",
        "   • En Obras Viales, todas las partidas deben analizarse sin excepción.",
        "",
        "4) Listado de Materiales (hoja 'Listado de Materiales')",
        "   • Incluye materiales, herramientas, máquinas y equipos de cada partida, en base a las EETT modificadas.",
        "",
        "5) Cotizaciones (hojas 'Cotización Materiales/Herramientas/Máquinas/Equipos')",
        "   • Deben ser del año en que se cursa la asignatura (no se aceptan cotizaciones de años anteriores).",
        "   • Materiales: 3 cotizaciones de 3 proveedores distintos (salvo exclusividad de proveedor único).",
        "   • Herramientas, máquinas y equipos: 1 cotización cada uno.",
        "   • Máquinas de construcción: informar precio de arriendo por hora o por día.",
        "   • Cada cotización debe respaldarse en PDF, nombrado con el material/ítem de referencia; "
        "si el respaldo agrupa varios materiales, nombrar el archivo 'VARIOS' y numerar correlativamente.",
        "",
        "Celdas en gris cursiva = ejemplos de la pauta. Reemplázalas por tus propios datos.",
    ]
    for i, texto in enumerate(filas, start=5):
        ws.cell(row=i, column=2, value=texto)
        ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        if texto.strip().startswith(("REGLAS", "1)", "2)", "3)", "4)", "5)")):
            ws.cell(row=i, column=2).font = FUENTE_SUBTITULO


def hoja_listado_actividades(wb) -> None:
    ws = wb.create_sheet("Listado de Actividades")
    ws.sheet_view.showGridLines = False
    _titulo(ws, "LISTADO DE ACTIVIDADES (ITEMIZADO)", 3)

    fila_encabezado = 3
    columnas = ["ITEM", "DESCRIPCIÓN / ACTIVIDAD", "UNIDAD"]
    anchos = [12, 70, 12]
    _encabezado(ws, fila_encabezado, columnas, anchos)

    ejemplos = [
        ["1.1", "Instalación de faenas", "gl"],
        ["1.2", "Cierros provisorios", "ml"],
        ["1.3", "Demoliciones y limpieza de terreno", "m2"],
        ["2.1", "Radier H-25 (e=10cm)", "m3"],
        ["2.2", "Porcelanato pavimentos exteriores Budnik (45x20cm)", "m2"],
        ["2.3", "Lana mineral AislánGlass (50mm)", "m2"],
    ]
    for i, valores in enumerate(ejemplos, start=fila_encabezado + 1):
        _fila_ejemplo(ws, i, valores)

    fila_fin = fila_encabezado + len(ejemplos) + FILAS_TRABAJO
    _bordear_area(ws, fila_encabezado + 1, fila_fin, 1, 3)

    dv_unidad = DataValidation(type="list", formula1=f'"{UNIDADES}"', allow_blank=True)
    ws.add_data_validation(dv_unidad)
    dv_unidad.add(f"C{fila_encabezado + 1}:C{fila_fin}")

    ws.cell(row=fila_encabezado + 1, column=2).comment = Comment(
        "El ITEM y la Descripción deben coincidir exactamente con la numeración y el "
        "nombre usado en las Especificaciones Técnicas modificadas.",
        "Pauta Proyecto de Título",
    )
    ws.auto_filter.ref = f"A{fila_encabezado}:C{fila_fin}"


def hoja_cubicaciones(wb) -> None:
    ws = wb.create_sheet("Cubicaciones")
    ws.sheet_view.showGridLines = False
    _titulo(ws, "CUBICACIONES", 6)

    fila_encabezado = 3
    columnas = [
        "ITEM",
        "DESCRIPCIÓN / ACTIVIDAD",
        "UNIDAD",
        "MEMORIA DE CÁLCULO (fórmula / desarrollo)",
        "CANTIDAD",
        "VERIFICACIÓN",
    ]
    anchos = [10, 46, 10, 46, 14, 26]
    _encabezado(ws, fila_encabezado, columnas, anchos)

    ejemplos = [
        [
            "2.1",
            "Radier H-25 (e=10cm)",
            "m3",
            "(12,35m x 8,20m) x 0,10m = 10,13 m3",
            10.13,
            None,
        ],
    ]
    for i, valores in enumerate(ejemplos, start=fila_encabezado + 1):
        _fila_ejemplo(ws, i, valores)

    fila_fin = fila_encabezado + len(ejemplos) + FILAS_TRABAJO
    for fila in range(fila_encabezado + 1, fila_fin + 1):
        ws.cell(row=fila, column=5).number_format = "0.00"
        formula = f'=IF(AND(C{fila}="un",E{fila}<>"",MOD(E{fila},1)<>0),"⚠ Debe ser entero (UN)","")'
        celda_verificacion = ws.cell(row=fila, column=6, value=formula)
        celda_verificacion.font = Font(color="C00000", italic=True)

    _bordear_area(ws, fila_encabezado + 1, fila_fin, 1, 6)

    dv_unidad = DataValidation(type="list", formula1=f'"{UNIDADES}"', allow_blank=True)
    ws.add_data_validation(dv_unidad)
    dv_unidad.add(f"C{fila_encabezado + 1}:C{fila_fin}")

    ws.cell(row=fila_encabezado + 1, column=4).comment = Comment(
        "Las cubicaciones son exactas: al menos 2 decimales, nunca redondear al entero "
        "superior. Registra siempre la fórmula/desarrollo que respalda el número.",
        "Pauta Proyecto de Título",
    )
    ws.auto_filter.ref = f"A{fila_encabezado}:F{fila_fin}"


def hoja_listado_materiales(wb) -> None:
    ws = wb.create_sheet("Listado de Materiales")
    ws.sheet_view.showGridLines = False
    _titulo(ws, "LISTADO DE MATERIALES, HERRAMIENTAS, MÁQUINAS Y EQUIPOS", 7)

    fila_encabezado = 3
    columnas = [
        "ITEM (partida)",
        "MATERIAL / HERRAMIENTA / MÁQUINA / EQUIPO",
        "TIPO",
        "CALIDAD / MATERIALIDAD / FORMATO / COLOR / MODELO / MARCA",
        "CANTIDAD REQUERIDA",
        "UNIDAD",
        "OBSERVACIONES",
    ]
    anchos = [12, 40, 14, 46, 16, 10, 30]
    _encabezado(ws, fila_encabezado, columnas, anchos)

    ejemplos = [
        [
            "1.3",
            "Pintura de alto tráfico fastrack blanco",
            "Material",
            "1 galón, línea alto tráfico, color blanco",
            2,
            "un",
            "",
        ],
        [
            "—",
            "Martillo carpintero 20 onzas",
            "Herramienta",
            "Stanley, 29mm",
            1,
            "un",
            "Uso transversal en varias partidas",
        ],
    ]
    for i, valores in enumerate(ejemplos, start=fila_encabezado + 1):
        _fila_ejemplo(ws, i, valores)

    fila_fin = fila_encabezado + len(ejemplos) + FILAS_TRABAJO
    _bordear_area(ws, fila_encabezado + 1, fila_fin, 1, 7)

    dv_tipo = DataValidation(type="list", formula1=f'"{TIPOS_MATERIAL}"', allow_blank=True)
    ws.add_data_validation(dv_tipo)
    dv_tipo.add(f"C{fila_encabezado + 1}:C{fila_fin}")

    dv_unidad = DataValidation(type="list", formula1=f'"{UNIDADES}"', allow_blank=True)
    ws.add_data_validation(dv_unidad)
    dv_unidad.add(f"F{fila_encabezado + 1}:F{fila_fin}")

    ws.auto_filter.ref = f"A{fila_encabezado}:G{fila_fin}"


def _hoja_cotizacion(
    wb,
    nombre_hoja: str,
    titulo_categoria: str,
    columna_item: str,
    columna_bien: str,
    proveedores: list[str],
    ejemplos: list[list],
    nota_precio: str | None = None,
) -> None:
    ws = wb.create_sheet(nombre_hoja)
    ws.sheet_view.showGridLines = False

    ncols = 2 + len(proveedores) + 1
    _titulo(
        ws,
        f'=CONCATENATE("{titulo_categoria}. PROYECTO ",Instrucciones!$B$2)',
        ncols,
    )

    fila_encabezado = 3
    columnas = [columna_item, columna_bien] + proveedores + ["NOMBRE ARCHIVO COTIZACIÓN"]
    anchos = [10, 42] + [16] * len(proveedores) + [26]
    _encabezado(ws, fila_encabezado, columnas, anchos)

    for i, valores in enumerate(ejemplos, start=fila_encabezado + 1):
        _fila_ejemplo(ws, i, valores)
        for col in range(3, 3 + len(proveedores)):
            ws.cell(row=i, column=col).number_format = '"$"#,##0'

    fila_fin = fila_encabezado + len(ejemplos) + FILAS_COTIZACION
    for fila in range(fila_encabezado + 1, fila_fin + 1):
        for col in range(3, 3 + len(proveedores)):
            ws.cell(row=fila, column=col).number_format = '"$"#,##0'

    _bordear_area(ws, fila_encabezado + 1, fila_fin, 1, ncols)

    if nota_precio:
        ws.cell(row=fila_encabezado + 1, column=3).comment = Comment(nota_precio, "Pauta Proyecto de Título")

    ws.auto_filter.ref = f"A{fila_encabezado}:{get_column_letter(ncols)}{fila_fin}"


def hoja_cotizacion_materiales(wb) -> None:
    ejemplos = [
        ["1", "Pintura de alto tráfico fastrack blanco 1 GL", 20480, 22320, 23090, "1.3 Pintura"],
        ["2", "1 1/2 x 5'' x 3.20m Pino dimensionado", 4025, 5079, 4354, "1.3 Pino"],
        ["3", "Terciado moldaje pino 15mm", 15160, 17870, 16899, "Varios01"],
    ]
    _hoja_cotizacion(
        wb,
        "Cotización Materiales",
        "COTIZACIÓN DE MATERIALES",
        "ITEM",
        "MATERIAL",
        ["PROVEEDOR 01", "PROVEEDOR 02", "PROVEEDOR 03"],
        ejemplos,
        nota_precio="Se requieren 3 cotizaciones de 3 proveedores distintos por material "
        "(salvo proveedor único/exclusivo, que se debe justificar).",
    )


def hoja_cotizacion_herramientas(wb) -> None:
    ejemplos = [
        ["1", "Martillo carpintero 20 onzas 29mm Stanley", 3085, "01Martillo"],
        ["2", "Chuzo hexagonal 1 x 1,7 metros negro Plasmet", 11280, "02Varios01"],
    ]
    _hoja_cotizacion(
        wb,
        "Cotización Herramientas",
        "COTIZACIÓN DE HERRAMIENTAS",
        "ITEM",
        "HERRAMIENTA",
        ["PROVEEDOR 01"],
        ejemplos,
        nota_precio="Herramientas: basta con 1 cotización.",
    )


def hoja_cotizacion_maquinas(wb) -> None:
    ejemplos = [
        ["1", "Retroexcavadora JCB motor 4.4 - 92HP (precio/hora)", 15000, "01Retro"],
        ["2", "Estación Total STONEX (precio/día)", 11200, "02EstacionTotal01"],
    ]
    _hoja_cotizacion(
        wb,
        "Cotización Máquinas",
        "COTIZACIÓN DE MÁQUINAS",
        "ITEM",
        "MÁQUINA (se moviliza mecánicamente)",
        ["PROVEEDOR (precio/hora o precio/día)"],
        ejemplos,
        nota_precio="Indicar arriendo por hora o por día. Considerar que normalmente incluye "
        "flete y operario especializado.",
    )


def hoja_cotizacion_equipos(wb) -> None:
    ejemplos = [
        ["1", "Betonera Rollmix 130lts", 249990, "01Betonera"],
        ["2", "Soldadora Ivertflash Katana 120", 141990, "02Soldadora"],
    ]
    _hoja_cotizacion(
        wb,
        "Cotización Equipos",
        "COTIZACIÓN DE EQUIPOS",
        "ITEM",
        "EQUIPO (no se moviliza mecánicamente)",
        ["PROVEEDOR"],
        ejemplos,
        nota_precio="Equipos: basta con 1 cotización.",
    )


def generar(path: str) -> None:
    wb = openpyxl.Workbook()
    hoja_instrucciones(wb)
    hoja_listado_actividades(wb)
    hoja_cubicaciones(wb)
    hoja_listado_materiales(wb)
    hoja_cotizacion_materiales(wb)
    hoja_cotizacion_herramientas(wb)
    hoja_cotizacion_maquinas(wb)
    hoja_cotizacion_equipos(wb)
    wb.save(path)


if __name__ == "__main__":
    destino = sys.argv[1] if len(sys.argv) > 1 else "Planilla_Formulacion_Proyecto_Entrega1.xlsx"
    generar(destino)
    print(f"Planilla generada en: {destino}")
