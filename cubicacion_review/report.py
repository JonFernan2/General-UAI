"""Generación del reporte de revisión (resumen en consola + Excel)."""
from __future__ import annotations

from collections import Counter

import openpyxl
from openpyxl.styles import Font, PatternFill

from .checks import Issue

ENCABEZADOS_DETALLE = ["Hoja", "Fila", "Código", "Categoría", "Severidad", "Mensaje"]

RELLENO_ERROR = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
RELLENO_ADVERTENCIA = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")


def resumen_texto(issues: list[Issue]) -> str:
    if not issues:
        return "No se encontraron observaciones. La cubicación pasó todas las validaciones."

    por_categoria = Counter(i.categoria for i in issues)
    errores = sum(1 for i in issues if i.severidad == "error")
    advertencias = sum(1 for i in issues if i.severidad == "advertencia")

    lineas = [f"Se encontraron {len(issues)} observaciones ({errores} errores, {advertencias} advertencias):"]
    for categoria, cantidad in sorted(por_categoria.items()):
        lineas.append(f"  - {categoria}: {cantidad}")
    return "\n".join(lineas)


def guardar_reporte_excel(issues: list[Issue], path: str) -> None:
    wb = openpyxl.Workbook()

    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    ws_resumen.append(["Total de observaciones", len(issues)])
    ws_resumen.append(["Errores", sum(1 for i in issues if i.severidad == "error")])
    ws_resumen.append(["Advertencias", sum(1 for i in issues if i.severidad == "advertencia")])
    ws_resumen.append([])
    ws_resumen.append(["Categoría", "Cantidad"])
    ws_resumen["A5"].font = Font(bold=True)
    ws_resumen["B5"].font = Font(bold=True)
    for categoria, cantidad in sorted(Counter(i.categoria for i in issues).items()):
        ws_resumen.append([categoria, cantidad])

    ws_detalle = wb.create_sheet("Detalle")
    ws_detalle.append(ENCABEZADOS_DETALLE)
    for celda in ws_detalle[1]:
        celda.font = Font(bold=True)

    for issue in issues:
        ws_detalle.append(
            [
                issue.hoja,
                issue.fila_excel,
                issue.codigo,
                issue.categoria,
                issue.severidad,
                issue.mensaje,
            ]
        )
        fila_actual = ws_detalle.max_row
        relleno = RELLENO_ERROR if issue.severidad == "error" else RELLENO_ADVERTENCIA
        for celda in ws_detalle[fila_actual]:
            celda.fill = relleno

    for columna, ancho in zip("ABCDEF", [20, 8, 15, 16, 12, 80]):
        ws_detalle.column_dimensions[columna].width = ancho

    wb.save(path)
