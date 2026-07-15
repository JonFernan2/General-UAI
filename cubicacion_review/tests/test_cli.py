import openpyxl

from cubicacion_review.cli import main

from .conftest import ENCABEZADO_ESTANDAR


def test_cli_extremo_a_extremo_sin_errores(crear_libro, tmp_path, capsys):
    path = crear_libro(
        "cub.xlsx",
        {"Hoja1": [ENCABEZADO_ESTANDAR, ["001", "Excavación", "m3", 10, 5, 50]]},
    )
    salida = str(tmp_path / "reporte.xlsx")

    codigo_salida = main([path, "--salida", salida])

    assert codigo_salida == 0
    capturado = capsys.readouterr()
    assert "No se encontraron observaciones" in capturado.out

    wb = openpyxl.load_workbook(salida)
    assert "Resumen" in wb.sheetnames
    assert "Detalle" in wb.sheetnames


def test_cli_reporta_errores_con_codigo_de_salida_1(crear_libro, tmp_path):
    path = crear_libro(
        "cub.xlsx",
        {"Hoja1": [ENCABEZADO_ESTANDAR, ["001", "Excavación", "m3", 10, 5, 999]]},
    )
    salida = str(tmp_path / "reporte.xlsx")

    codigo_salida = main([path, "--salida", salida])

    assert codigo_salida == 1
