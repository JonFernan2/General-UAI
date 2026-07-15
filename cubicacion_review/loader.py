"""Carga de archivos Excel de cubicación y detección de encabezados/columnas."""
from __future__ import annotations

from dataclasses import dataclass, field

import openpyxl

from .config import REQUIRED_COLUMNS, Config, normalizar_texto

MAX_FILAS_BUSQUEDA_ENCABEZADO = 20


@dataclass
class Fila:
    fila_excel: int  # número de fila tal como aparece en Excel (1-based)
    valores: dict[str, object]  # nombre canónico de columna -> valor de la celda


@dataclass
class Hoja:
    nombre: str
    columnas: dict[str, int] = field(default_factory=dict)  # canónico -> índice 0-based
    columnas_faltantes: list[str] = field(default_factory=list)
    filas: list[Fila] = field(default_factory=list)
    encabezado_encontrado: bool = False


def _detectar_encabezado(ws, alias_map: dict[str, str]) -> tuple[int, dict[str, int]] | None:
    """Busca, en las primeras filas, la fila que mejor calza con los alias de columna.

    Devuelve (índice de fila 1-based, {canonico: índice de columna 0-based}).
    """
    mejor: tuple[int, dict[str, int]] | None = None
    max_matches = 0
    limite = min(ws.max_row, MAX_FILAS_BUSQUEDA_ENCABEZADO)
    for fila_idx in range(1, limite + 1):
        columnas: dict[str, int] = {}
        for col_idx, celda in enumerate(ws[fila_idx]):
            canonico = alias_map.get(normalizar_texto(celda.value))
            if canonico and canonico not in columnas:
                columnas[canonico] = col_idx
        if len(columnas) > max_matches:
            max_matches = len(columnas)
            mejor = (fila_idx, columnas)
    if mejor is None or max_matches < 3:
        return None
    return mejor


def _fila_vacia(valores: dict[str, object]) -> bool:
    return all(v is None or str(v).strip() == "" for v in valores.values())


def cargar_libro(path: str, config: Config) -> list[Hoja]:
    wb = openpyxl.load_workbook(path, data_only=True)
    alias_map = config.alias_normalizados()

    hojas: list[Hoja] = []
    for nombre_hoja in wb.sheetnames:
        if nombre_hoja in config.hojas_ignoradas:
            continue
        ws = wb[nombre_hoja]
        hoja = Hoja(nombre=nombre_hoja)

        encabezado = _detectar_encabezado(ws, alias_map)
        if encabezado is None:
            hoja.columnas_faltantes = list(REQUIRED_COLUMNS)
            hojas.append(hoja)
            continue

        fila_encabezado, columnas = encabezado
        hoja.columnas = columnas
        hoja.encabezado_encontrado = True
        hoja.columnas_faltantes = [c for c in REQUIRED_COLUMNS if c not in columnas]

        for fila_idx in range(fila_encabezado + 1, ws.max_row + 1):
            celdas = ws[fila_idx]
            valores = {
                canonico: celdas[col_idx].value if col_idx < len(celdas) else None
                for canonico, col_idx in columnas.items()
            }
            if _fila_vacia(valores):
                continue
            hoja.filas.append(Fila(fila_excel=fila_idx, valores=valores))

        hojas.append(hoja)

    return hojas
